import openpyxl
from openpyxl.chart import BarChart, Series, PieChart,Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Font
from openpyxl.chart.series import DataPoint


def generate_charts(filename):
    wb = openpyxl.load_workbook(filename)
    mysheet = wb.get_active_sheet()

    # 找到最大行
    for i in range(3,30):
        if mysheet['B'+str(i)].value == None:
            max_row = i - 1
            break

    # 
    ft = Font(bold=True)
    mysheet['B' + str(max_row+1)].value = 'total'
    mysheet['B' + str(max_row+1)].font = ft

    k = 54
    t = 61
    for x in 'CDEFGHIJKLMN':
        x_max = x + str(max_row)
        x_cell = x + str(max_row+1)
        mysheet[x_cell].value = '=SUM(' + x + str(4) + ':' + x_max+ ')'
        
        if x >'D' and x < 'J':
            mysheet['C'+str(k)].value = mysheet[x_cell].value
            k += 1
        if x >= 'J' and x <= 'N':
            mysheet['C'+str(t)].value = mysheet[x_cell].value
            t += 1

        mysheet[x_cell].font = ft


    chart = BarChart()
    chart.style = 11
    chart.type = "bar"
    chart.title = "问题燃尽情况"
    # 系列重叠
    chart.overlap = 100
    # 添加数据标签
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showVal = True


    data = Reference(mysheet, min_col=3, min_row=3, max_row=max_row, max_col=4)
    cats = Reference(mysheet, min_col=2, min_row=4, max_row=max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.shape = 4
    mysheet.add_chart(chart, 'B'+str(max_row+3))


    pie1 = PieChart()
    labels = Reference(mysheet, min_col=2, min_row=54, max_row=58)
    data = Reference(mysheet, min_col=3, min_row=53, max_row=58)
    pie1.add_data(data, titles_from_data=True)
    pie1.set_categories(labels)
    pie1.title = "问题分类占比"
    pie1.dataLabels = DataLabelList()
    pie1.dataLabels.showVal = True

    # Cut the first slice out of the pie
    slice = DataPoint(idx=0, explosion=20)
    pie1.series[0].data_points = [slice]

    mysheet.add_chart(pie1,'I'+str(max_row+3))

    pie2 = PieChart()
    labels = Reference(mysheet, min_col=2, min_row=61, max_row=65)
    data = Reference(mysheet, min_col=3, min_row=60, max_row=65)
    pie2.add_data(data, titles_from_data=True)
    pie2.set_categories(labels)
    pie2.title = "严重级别占比"
    pie2.dataLabels = DataLabelList()
    pie2.dataLabels.showVal = True

    # Cut the first slice out of the pie
    slice = DataPoint(idx=0, explosion=20)
    pie2.series[0].data_points = [slice]

    mysheet.add_chart(pie2,'B'+str(max_row+19))

    wb.save(filename)
